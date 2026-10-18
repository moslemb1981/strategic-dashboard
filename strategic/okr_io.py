# -*- coding: utf-8 -*-
"""ورود/خروج اکسل OKR — مشترک بین دستور import_okr و دکمه‌های صفحه OKR.

قالب فایل (همان فایل OKR نهایی): سه شیت «OKR شرکت»، «OKR معاونت‌ها»، «OKR مدیریت‌ها».
ستون‌ها: ۱ کد، ۲ واحد، ۳ شرح هدف/نتیجه کلیدی، ۴ هدف بالادستی، ۵ کد شاخص، ۶ واحد اندازه‌گیری،
۷ جهت بهبود (↓ کمتر بهتر). ستون ۸ اگر عنوانش «وزن» باشد وزن نتیجه کلیدی خوانده می‌شود
(فایل‌های خروجی همین صفحه این ستون را دارند؛ در فایل اصلی نیست و وزن ۱۰۰ می‌شود)."""
import re

from django.db import transaction

from .models import CompanyKPI, OKRKeyResult, OKRObjective, OperationalKPI, OrgUnit

SHEETS = (("OKR شرکت", "company"), ("OKR معاونت‌ها", "deputy"), ("OKR مدیریت‌ها", "management"))
KR_RE = re.compile(r"-K\d+$")
HEAD = ["کد", "واحد", "شرح هدف (Objective) / نتیجه‌ی کلیدی (KR)", "هدف بالادستی (هم‌راستایی)", "کد شاخص",
        "واحد اندازه‌گیری", "جهت بهبود", "وزن", "هدف ۱۴۰۵ (سیستم)", "عملکرد ۱۴۰۵ (سیستم)", "٪ پیشرفت"]


def _norm(s):
    return re.sub(r"\s+", " ", str(s or "").replace("‌", " ").replace("ي", "ی").replace("ك", "ک")).strip()


def _w(v):
    try:
        return max(0, min(int(float(v)), 100))
    except (TypeError, ValueError):
        return 100


class OKRImportError(Exception):
    pass


def import_workbook(wb, keep=False, dry_run=False):
    for n, _ in SHEETS:
        if n not in wb.sheetnames:
            raise OKRImportError(f"شیت «{n}» در فایل نیست.")
    op = {k.code: k for k in OperationalKPI.objects.all()}
    macro = {k.code: k for k in CompanyKPI.objects.all()}
    units = {_norm(u.name): u for u in OrgUnit.objects.all()}

    def unit_of(name):
        u = units.get(_norm(name))
        if not u:
            u = OrgUnit.objects.create(
                name=str(name).strip(), kind="ceo_direct" if "مستقل" in str(name) else "deputy", order=len(units))
            units[_norm(name)] = u
        return u

    stats = {"obj": 0, "kr": 0, "nokpi": 0}
    with transaction.atomic():
        if not keep:
            OKRObjective.objects.all().delete()
        made, pending = {}, []
        for sheet, level in SHEETS:
            ws = wb[sheet]
            has_w = str(ws.cell(row=4, column=8).value or "").strip() == "وزن"
            grp, cur, order = None, None, 0
            for r in ws.iter_rows(min_row=5, values_only=True):
                if r[0] and all(x is None for x in r[1:]):
                    grp = str(r[0])
                    continue
                if not r[0]:
                    continue
                code = str(r[0]).strip()
                if KR_RE.search(code):
                    if cur is None:
                        continue
                    kc = str(r[4]).strip() if r[4] not in (None, "—", "") else None
                    k_op = op.get(kc)
                    k_macro = macro.get(kc) if kc and not k_op else None
                    OKRKeyResult.objects.create(
                        objective=cur, code=code, title=str(r[2]).strip(), operational_kpi=k_op, company_kpi=k_macro,
                        lower_is_better="↓" in str(r[6] or ""), weight=_w(r[7]) if has_w and len(r) > 7 and r[7] is not None else 100,
                        order=cur.key_results.count() + 1)
                    stats["kr"] += 1
                    if not (k_op or k_macro):
                        stats["nokpi"] += 1
                    continue
                order += 1
                unit, mgmt = None, ""
                if level == "deputy":
                    unit = unit_of(r[1])
                elif level == "management":
                    mgmt = str(r[1]).strip()
                    unit = unit_of((grp or "").split("|")[-1].strip())
                par = (str(r[3] or "").split(" — ")[0]).strip()
                if OKRObjective.objects.filter(code=code).exists():
                    raise OKRImportError(f"کد هدف «{code}» تکراری است (برای افزودن بدون پاک‌کردن، کدها باید تازه باشند).")
                cur = OKRObjective.objects.create(
                    level=level, code=code, title=re.sub(r"^هدف:\s*", "", str(r[2])).strip(),
                    org_unit=unit, management=mgmt, order=order)
                made[code] = cur
                pending.append((cur, par))
                stats["obj"] += 1
        for ob, par in pending:
            if par in made:
                ob.parent = made[par]
            elif par:
                ob.parent_text = par
            else:
                continue
            ob.save(update_fields=["parent", "parent_text"])
        if dry_run:
            transaction.set_rollback(True)
    return stats


def export_workbook():
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    fill = PatternFill(start_color="1B2430", end_color="1B2430", fill_type="solid")
    ofill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
    widths = [12, 26, 60, 40, 12, 14, 12, 8, 14, 14, 10]
    for sheet, level in SHEETS:
        ws = wb.create_sheet(sheet)
        ws.sheet_view.rightToLeft = True
        ws.cell(row=1, column=1, value=sheet).font = Font(bold=True, size=14)
        for c, h in enumerate(HEAD, 1):
            cell = ws.cell(row=4, column=c, value=h)
            cell.fill, cell.font = fill, Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for i, wd in enumerate(widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = wd
        row, last_grp = 5, None
        qs = (OKRObjective.objects.filter(level=level).select_related("org_unit", "parent")
              .prefetch_related("key_results__operational_kpi", "key_results__company_kpi")
              .order_by("org_unit__kind", "org_unit__order", "management", "order", "code"))
        for o in qs:
            if level == "management":
                grp = f"{o.management}   |   {o.org_unit.name if o.org_unit else ''}"
            elif level == "deputy":
                grp = o.org_unit.name if o.org_unit else ""
            else:
                grp = None
            if grp and grp != last_grp:
                ws.cell(row=row, column=1, value=grp).font = Font(bold=True)
                row += 1
                last_grp = grp
            unit_txt = o.management if level == "management" else (o.org_unit.name if o.org_unit else "شرکت")
            par = (f"{o.parent.code} — {o.parent.title}" if o.parent else (o.parent_text or ""))
            vals = [o.code, unit_txt, "هدف: " + o.title, par]
            for c, v in enumerate(vals, 1):
                cell = ws.cell(row=row, column=c, value=v)
                cell.fill, cell.font = ofill, Font(bold=True)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            for c in range(5, 12):
                ws.cell(row=row, column=c).fill = ofill
            row += 1
            for kr in o.key_results.all():
                k = kr.kpi
                vals = [kr.code, unit_txt, kr.title, None, kr.kpi_code or "—", (k.unit if k else ""),
                        "↓ کاهشی" if kr.lower_is_better else "↑ افزایشی", kr.weight, kr.target, kr.actual, kr.pct]
                for c, v in enumerate(vals, 1):
                    cell = ws.cell(row=row, column=c, value=v)
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                row += 1
        ws.freeze_panes = "A5"
    return wb
